import csv
import hmac
import io
import json
import math
import os
import secrets
import sqlite3
import threading
import time as clock
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from urllib.parse import urlsplit

from flask import (
    Flask, abort, flash, g, jsonify, make_response, redirect, render_template,
    request, send_file, session, url_for,
)

from Services.hotspot_forecasting import (
    forecast_hotspots,
    recommend_safety_actions,
)
from Services.incident_services import IncidentService
from Services.public_signals import fetch_public_safety_signals
from Services.risk_assessment import RiskAssessment
from storage.database_manager import DatabaseManager


PROJECT_DIR = Path(__file__).resolve().parent
GEOCODE_LOCK = threading.Lock()
LAST_GEOCODE_REQUEST = 0.0
FORM_FIELDS = (
    ("cause_code", "Incident reference"),
    ("date", "Date"),
    ("time", "Time"),
    ("county", "County"),
    ("road", "Road"),
    ("place", "Location"),
    ("base", "Nearest base"),
    ("accident_info", "Accident details"),
    ("victim", "Road user affected"),
)


def _form_date(value):
    for date_format in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, date_format).date().isoformat()
        except ValueError:
            continue
    return ""


def _form_time(value):
    if not value:
        return ""
    if value.isdigit() and len(value) == 3:
        value = f"0{value}"
    for pattern in ("%H:%M", "%H%M", "%H%M%S"):
        try:
            return datetime.strptime(value, pattern).strftime("%H:%M")
        except ValueError:
            continue
    return ""


def _risk(incident, assessor):
    score = assessor.calculate_risk_score(incident)
    return {"score": score, "level": assessor.get_risk_level(score)}


def _hotspots(incidents):
    counts = Counter(
        incident["place"].strip()
        for incident in incidents
        if incident.get("place", "").strip()
    )
    fatality_counts = Counter(
        incident["place"].strip()
        for incident in incidents
        if incident.get("place", "").strip()
        and int(incident.get("fatalities") or 0) > 0
    )
    return [
        {
            "location": location,
            "accident_count": count,
            "fatal_incidents": fatality_counts[location],
        }
        for location, count in counts.most_common()
    ]


def _risk_summary(incidents, assessor):
    summary = Counter(
        _risk(incident, assessor)["level"] for incident in incidents
    )
    return {
        level: summary[level]
        for level in ("Critical", "High", "Moderate", "Low")
    }


def _parse_date(value):
    for date_format in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, date_format).date()
        except (TypeError, ValueError):
            continue
    return None


def _filter_incidents(incidents, assessor):
    counties = sorted({
        item.get("county", "").strip()
        for item in incidents if item.get("county", "").strip()
    })
    roads = sorted({
        item.get("road", "").strip()
        for item in incidents if item.get("road", "").strip()
    })
    filters = {
        "county": request.args.get("county", "").strip(),
        "road": request.args.get("road", "").strip(),
        "start_date": request.args.get("start_date", "").strip(),
        "end_date": request.args.get("end_date", "").strip(),
        "severity": request.args.get("severity", "").strip(),
    }
    start_date = _parse_date(filters["start_date"])
    end_date = _parse_date(filters["end_date"])
    if filters["start_date"] and not start_date:
        abort(400, description="Start date must be a valid date.")
    if filters["end_date"] and not end_date:
        abort(400, description="End date must be a valid date.")
    if start_date and end_date and start_date > end_date:
        abort(400, description="Start date must be on or before end date.")
    valid_levels = {"Low", "Moderate", "High", "Critical"}
    if filters["severity"] and filters["severity"] not in valid_levels:
        abort(400, description="Choose a valid risk category.")

    filtered = []
    for incident in incidents:
        incident_date = _parse_date(incident.get("date", ""))
        risk_level = assessor.get_risk_level(
            assessor.calculate_risk_score(incident)
        )
        if filters["county"] and incident.get("county", "").strip() != filters["county"]:
            continue
        if filters["road"] and incident.get("road", "").strip() != filters["road"]:
            continue
        if start_date and (not incident_date or incident_date < start_date):
            continue
        if end_date and (not incident_date or incident_date > end_date):
            continue
        if filters["severity"] and risk_level != filters["severity"]:
            continue
        filtered.append(incident)
    return filtered, filters, counties, roads


def _chart_data(incidents):
    monthly_counts = Counter()
    for incident in incidents:
        incident_date = _parse_date(incident.get("date", ""))
        if incident_date:
            monthly_counts[incident_date.strftime("%Y-%m")] += 1
    months = sorted(monthly_counts)
    hotspots = _hotspots(incidents)[:8]
    return {
        "monthly": {
            "labels": months,
            "values": [monthly_counts[month] for month in months],
        },
        "hotspots": {
            "labels": [item["location"] for item in hotspots],
            "values": [item["accident_count"] for item in hotspots],
        },
    }


def create_app(data_file=None, admin_password=None, database_path=None):
    app = Flask(__name__)
    app.secret_key = os.environ.get("FLASK_SECRET_KEY") or secrets.token_hex(32)
    configured_admin_password = (
        admin_password or os.environ.get("ROAD_SAFETY_ADMIN_PASSWORD")
    )
    csv_path = Path(data_file or PROJECT_DIR / "incidents.csv")
    database_path = Path(
        database_path
        or os.environ.get("TUSALAMA_DATABASE_PATH")
        or PROJECT_DIR / "data" / "tusalama.sqlite3"
    )
    app.config.update(
        ADMIN_PASSWORD_CONFIGURED=bool(configured_admin_password),
        DATABASE_PATH=database_path,
        STATIC_ASSET_VERSION=str(max(
            (PROJECT_DIR / "static" / filename).stat().st_mtime_ns
            for filename in ("dashboard.js", "styles.css")
        )),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("FLASK_COOKIE_SECURE") == "1",
    )

    database = DatabaseManager(database_path)
    app.config["CSV_MIGRATION"] = database.migrate_csv(csv_path)
    database.ensure_admin(
        os.environ.get("ROAD_SAFETY_ADMIN_USERNAME", "admin"),
        os.environ.get("ROAD_SAFETY_ADMIN_EMAIL", "admin@tusalama.local"),
        configured_admin_password,
    )
    incident_service = IncidentService(database)
    risk_assessment = RiskAssessment()

    def current_scope():
        user = g.current_user
        if user["role"] == "road_user":
            return {"owner_user_id": user["id"]}
        if user["role"] == "agency":
            return {"agency_name": user["agency_name"]}
        return {}

    def visible_incidents():
        return incident_service.view_all_incidents(**current_scope())

    public_signals_cache = {"fetched_at": 0.0, "payload": None}

    @app.before_request
    def protect_admin_area():
        if request.endpoint == "static":
            return None
        session.setdefault("csrf_token", secrets.token_urlsafe(32))
        g.current_user = database.get_user(session.get("user_id"))
        if request.method == "POST":
            token = request.form.get("csrf_token", "")
            if not hmac.compare_digest(token, session["csrf_token"]):
                abort(400, description="Invalid form security token.")
        if request.endpoint in {"login", "signup", "setup_admin"}:
            return None
        if not g.current_user:
            return redirect(url_for("login"))
        if (
            g.current_user["role"] == "agency"
            and g.current_user["status"] == "pending"
            and request.endpoint not in {"agency_pending", "profile", "logout"}
        ):
            return redirect(url_for("agency_pending"))
        if (
            request.endpoint in {
                "admin", "approve_agency", "database_backup",
            }
            and g.current_user["role"] != "admin"
        ):
            abort(403)
        if (
            request.endpoint in {
                "analytics", "intelligence", "public_signals",
            }
            and g.current_user["role"] == "road_user"
        ):
            abort(403, description="Hotspot analytics are available to agencies and administrators.")
        return None

    @app.context_processor
    def inject_template_security():
        return {
            "csrf_token": session.get("csrf_token", ""),
            "current_user": getattr(g, "current_user", None),
            "static_asset_version": app.config["STATIC_ASSET_VERSION"],
        }

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            account = database.authenticate(
                request.form.get("username", ""),
                request.form.get("password", ""),
            )
            if not account:
                flash("Username/email or password is incorrect.", "error")
                return redirect(url_for("login"))
            session.clear()
            session["user_id"] = account["id"]
            session["csrf_token"] = secrets.token_urlsafe(32)
            if account["role"] == "agency" and account["status"] == "pending":
                return redirect(url_for("agency_pending"))
            destination = request.args.get("next", "")
            parsed_destination = urlsplit(destination)
            if (
                not destination.startswith("/")
                or destination.startswith("//")
                or "\\" in destination
                or parsed_destination.netloc
                or parsed_destination.scheme
            ):
                destination = url_for("dashboard")
            return redirect(destination)
        return render_template(
            "login.html",
            password_configured=(
                app.config["ADMIN_PASSWORD_CONFIGURED"]
                or not database.can_bootstrap_admin()
            ),
            admin_setup_available=(
                not app.config["ADMIN_PASSWORD_CONFIGURED"]
                and database.can_bootstrap_admin()
            ),
        )

    @app.route("/setup-admin", methods=["GET", "POST"])
    def setup_admin():
        if (
            app.config["ADMIN_PASSWORD_CONFIGURED"]
            or not database.can_bootstrap_admin()
        ):
            abort(404)
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            if (
                len(username) < 3
                or len(username) > 40
                or not all(
                    character.isalnum() or character in "._-"
                    for character in username
                )
            ):
                flash("Username must be 3–40 characters using letters, numbers, dot, dash, or underscore.", "error")
                return render_template("setup_admin.html"), 400
            if (
                len(email) > 254
                or email.count("@") != 1
                or "." not in email.rsplit("@", 1)[1]
            ):
                flash("Enter a valid email address.", "error")
                return render_template("setup_admin.html"), 400
            if len(password) < 12:
                flash("Choose a password with at least 12 characters.", "error")
                return render_template("setup_admin.html"), 400
            if password != request.form.get("confirm_password", ""):
                flash("The passwords do not match.", "error")
                return render_template("setup_admin.html"), 400
            try:
                admin_id = database.create_first_admin(
                    username, email, password
                )
            except sqlite3.IntegrityError:
                flash("That username or email is already registered.", "error")
                return render_template("setup_admin.html"), 400
            if admin_id is None:
                abort(404)
            session.clear()
            session["user_id"] = admin_id
            session["csrf_token"] = secrets.token_urlsafe(32)
            flash("Administrator account created. Welcome to Tusalama.", "success")
            return redirect(url_for("dashboard"))
        return render_template("setup_admin.html")

    @app.route("/signup", methods=["GET", "POST"])
    def signup():
        if (
            not app.config["ADMIN_PASSWORD_CONFIGURED"]
            and database.can_bootstrap_admin()
        ):
            return redirect(url_for("setup_admin"))
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            role = request.form.get("role", "")
            agency_name = request.form.get("agency_name", "").strip()
            if (
                len(username) < 3
                or len(username) > 40
                or not all(
                    character.isalnum() or character in "._-"
                    for character in username
                )
            ):
                flash("Username must be 3–40 characters using letters, numbers, dot, dash, or underscore.", "error")
                return render_template("signup.html"), 400
            if (
                len(email) > 254
                or email.count("@") != 1
                or "." not in email.rsplit("@", 1)[1]
            ):
                flash("Enter a valid email address.", "error")
                return render_template("signup.html"), 400
            if len(password) < 12:
                flash("Choose a password with at least 12 characters.", "error")
                return render_template("signup.html"), 400
            if role not in {"road_user", "agency"}:
                flash("Choose a road user or road safety agency account.", "error")
                return render_template("signup.html"), 400
            if role == "agency" and len(agency_name) < 3:
                flash("Enter the name of your road safety agency.", "error")
                return render_template("signup.html"), 400
            try:
                database.create_user(
                    username, email, password, role, agency_name
                )
            except sqlite3.IntegrityError:
                flash("That username or email is already registered.", "error")
                return render_template("signup.html"), 400
            if role == "agency":
                flash("Agency registration submitted for administrator approval.", "success")
            else:
                flash("Your road user account is ready. Please sign in.", "success")
            return redirect(url_for("login"))
        return render_template("signup.html")

    @app.get("/agency-pending")
    def agency_pending():
        return render_template("agency_pending.html")

    @app.get("/profile")
    def profile():
        return render_template("profile.html", profile=g.current_user)

    @app.post("/admin/agencies/<int:user_id>/approve")
    def approve_agency(user_id):
        if database.approve_agency(user_id):
            flash("Agency account approved.", "success")
        else:
            flash("Pending agency account not found.", "error")
        return redirect(url_for("admin"))

    @app.get("/admin/database-backup")
    def database_backup():
        backup_dir = database.database_path.parent / "backups"
        backup_path = backup_dir / (
            f"tusalama-backup-{datetime.now().strftime('%Y%m%d-%H%M%S')}.sqlite3"
        )
        database.backup(backup_path)
        return send_file(
            backup_path,
            as_attachment=True,
            download_name=backup_path.name,
            mimetype="application/vnd.sqlite3",
        )

    @app.get("/api/geocode")
    def geocode():
        query = request.args.get("q", "").strip()
        if len(query) < 3 or len(query) > 180:
            abort(400, description="Enter a location with 3–180 characters.")
        global LAST_GEOCODE_REQUEST
        with GEOCODE_LOCK:
            wait_seconds = 1 - (clock.monotonic() - LAST_GEOCODE_REQUEST)
            if wait_seconds > 0:
                clock.sleep(wait_seconds)
            LAST_GEOCODE_REQUEST = clock.monotonic()

        parameters = urllib.parse.urlencode({
            "q": f"{query}, Kenya",
            "format": "jsonv2",
            "countrycodes": "ke",
            "limit": "1",
        })
        request_url = (
            "https://nominatim.openstreetmap.org/search?" + parameters
        )
        outbound_request = urllib.request.Request(
            request_url,
            headers={
                "User-Agent": "TusalamaRoadSafety/1.0 (location lookup)",
                "Accept": "application/json",
            },
        )
        try:
            with urllib.request.urlopen(outbound_request, timeout=8) as result:
                candidates = json.load(result)
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as error:
            app.logger.warning("OpenStreetMap location lookup failed: %s", error)
            return jsonify(error="Location search is temporarily unavailable."), 502
        if not candidates:
            return jsonify(error="No matching Kenyan location was found."), 404
        return jsonify({
            "latitude": float(candidates[0]["lat"]),
            "longitude": float(candidates[0]["lon"]),
            "display_name": candidates[0].get("display_name", ""),
        })

    @app.post("/logout")
    def logout():
        session.clear()
        flash("You have been signed out.", "success")
        return redirect(url_for("login"))

    @app.get("/admin")
    def admin():
        return render_template(
            "admin.html",
            incident_count=database.count_incidents(),
            data_file=database.database_path.name,
            pending_agencies=database.pending_agencies(),
            migration=app.config["CSV_MIGRATION"],
            social_sources_configured=bool(
                os.environ.get("TUSALAMA_BSKY_HANDLES")
                or os.environ.get("TUSALAMA_X_USER_IDS")
            ),
        )

    @app.get("/")
    def dashboard():
        all_incidents = visible_incidents()
        incidents, filters, counties, roads = _filter_incidents(
            all_incidents, risk_assessment
        )
        fatalities = sum(int(item.get("fatalities") or 0) for item in incidents)
        hotspots = _hotspots(incidents)
        return render_template(
            "dashboard.html",
            today=datetime.now().strftime("%A, %d %B %Y").upper(),
            incidents=incidents,
            total_incidents=len(incidents),
            total_fatalities=fatalities,
            hotspot_count=sum(item["accident_count"] >= 3 for item in hotspots),
            hotspots=hotspots[:5],
            risk_summary=_risk_summary(incidents, risk_assessment),
            filters=filters,
            counties=counties,
            roads=roads,
            chart_data=_chart_data(incidents),
        )

    @app.get("/api/chart-data")
    def chart_data():
        incidents, _, _, _ = _filter_incidents(
            visible_incidents(), risk_assessment
        )
        return jsonify(_chart_data(incidents))

    @app.get("/incidents")
    def incidents():
        all_incidents = visible_incidents()
        all_incidents, filters, counties, roads = _filter_incidents(
            all_incidents, risk_assessment
        )
        query = request.args.get("q", "").strip().casefold()
        filtered = [
            incident for incident in all_incidents
            if not query or any(
                query in str(value).casefold() for value in incident.values()
            )
        ]
        filtered.sort(
            key=lambda incident: _parse_date(incident.get("date", "")) or date.min,
            reverse=True,
        )
        page_size = 50
        page_count = max(1, math.ceil(len(filtered) / page_size))
        page = max(1, request.args.get("page", 1, type=int) or 1)
        page = min(page, page_count)
        page_start = (page - 1) * page_size
        return render_template(
            "incidents.html",
            incidents=filtered[page_start:page_start + page_size],
            total_records=len(filtered),
            page=page,
            page_size=page_size,
            page_count=page_count,
            query=request.args.get("q", ""),
            assessor=risk_assessment,
            filters=filters,
            counties=counties,
            roads=roads,
        )

    def _incident_form(incident=None, original_cause_code=None):
        values = dict(incident or {})
        if request.method == "GET":
            values["date"] = _form_date(values.get("date", ""))
            values["time"] = _form_time(values.get("time", ""))
            return render_template(
                "incident_form.html",
                incident=values,
                fields=FORM_FIELDS,
                editing=incident is not None,
            )

        submitted = {
            key: request.form.get(key, "").strip() for key, _ in FORM_FIELDS
        }
        submitted["fatalities"] = request.form.get("fatalities", "").strip()
        submitted["latitude"] = request.form.get("latitude", "").strip()
        submitted["longitude"] = request.form.get("longitude", "").strip()
        errors = []
        if not all(submitted.get(field, "") for field, _ in FORM_FIELDS):
            errors.append("Complete all required fields.")
        if not submitted["fatalities"].isdigit():
            errors.append("Fatalities must be a whole number of zero or more.")
        if submitted["date"]:
            try:
                date.fromisoformat(submitted["date"])
            except ValueError:
                errors.append("Enter a valid accident date.")
        if submitted["time"]:
            try:
                time.fromisoformat(submitted["time"])
            except ValueError:
                errors.append("Enter a valid accident time.")
        has_latitude = bool(submitted["latitude"])
        has_longitude = bool(submitted["longitude"])
        if has_latitude != has_longitude:
            errors.append("Enter both latitude and longitude, or leave both blank.")
        elif has_latitude:
            try:
                latitude = float(submitted["latitude"])
                longitude = float(submitted["longitude"])
                if not -90 <= latitude <= 90 or not -180 <= longitude <= 180:
                    raise ValueError
                submitted["latitude"] = str(latitude)
                submitted["longitude"] = str(longitude)
            except ValueError:
                errors.append("Enter valid coordinates (latitude -90 to 90, longitude -180 to 180).")

        duplicate = incident_service.search_incident(
            submitted["cause_code"], **current_scope()
        )
        if duplicate and submitted["cause_code"] != original_cause_code:
            errors.append("That incident reference is already in use.")

        if errors:
            for error in errors:
                flash(error, "error")
            return render_template(
                "incident_form.html",
                incident=submitted,
                fields=FORM_FIELDS,
                editing=incident is not None,
            ), 400

        submitted["time"] = submitted["time"].replace(":", "")
        if incident is None:
            result = incident_service.register_incident(
                submitted, reporter_id=g.current_user["id"]
            )
        else:
            updates = {
                key: value for key, value in submitted.items()
                if key != "cause_code"
            }
            if submitted["cause_code"] != original_cause_code:
                updates["cause_code"] = submitted["cause_code"]
            result = incident_service.update_incident(
                original_cause_code, updates, **current_scope()
            )
        flash(result, "success")
        return redirect(url_for("incidents"))

    @app.route("/incidents/new", methods=["GET", "POST"])
    def new_incident():
        return _incident_form()

    @app.route("/incidents/<cause_code>/edit", methods=["GET", "POST"])
    def edit_incident(cause_code):
        incident = incident_service.search_incident(
            cause_code, **current_scope()
        )
        if incident is None:
            flash("Incident not found.", "error")
            return redirect(url_for("incidents"))
        return _incident_form(incident, cause_code)

    @app.post("/incidents/<cause_code>/delete")
    def delete_incident(cause_code):
        flash(
            incident_service.delete_incident(
                cause_code, **current_scope()
            ),
            "success",
        )
        return redirect(url_for("incidents"))

    @app.get("/incidents/<cause_code>/risk")
    def assess_incident(cause_code):
        incident = incident_service.search_incident(
            cause_code, **current_scope()
        )
        if incident is None:
            flash("Incident not found.", "error")
            return redirect(url_for("incidents"))
        score = risk_assessment.calculate_risk_score(incident)
        level = risk_assessment.get_risk_level(score)
        accident_score = risk_assessment.assess_accident_type(
            incident.get("accident_info", "")
        )
        fatality_score = risk_assessment.assess_fatalities(
            incident.get("fatalities", 0)
        )
        return render_template(
            "risk.html",
            incident=incident,
            score=score,
            level=level,
            accident_score=accident_score,
            fatality_score=fatality_score,
        )

    @app.get("/analytics")
    def analytics():
        all_incidents = visible_incidents()
        incidents, filters, counties, roads = _filter_incidents(
            all_incidents, risk_assessment
        )
        return render_template(
            "analytics.html",
            hotspots=_hotspots(incidents),
            road_counts=Counter(
                incident.get("road", "").strip()
                for incident in incidents
                if incident.get("road", "").strip()
            ).most_common(),
            risk_summary=_risk_summary(incidents, risk_assessment),
            accident_types=Counter(
                incident.get("accident_info", "Unspecified").strip().upper()
                for incident in incidents
            ).most_common(6),
            incidents=incidents,
            filters=filters,
            counties=counties,
            roads=roads,
            chart_data=_chart_data(incidents),
            map_points=[
                {
                    "latitude": float(item["latitude"]),
                    "longitude": float(item["longitude"]),
                    "location": item.get("place", ""),
                    "road": item.get("road", ""),
                    "county": item.get("county", ""),
                    "fatalities": int(item.get("fatalities") or 0),
                    "count": 1,
                }
                for item in incidents
                if item.get("latitude", "").strip()
                and item.get("longitude", "").strip()
            ],
        )

    @app.get("/intelligence")
    def intelligence():
        incidents = visible_incidents()
        return render_template(
            "intelligence.html",
            forecast=forecast_hotspots(incidents),
            recommendations=recommend_safety_actions(incidents),
            source_configuration={
                "x": bool(
                    os.environ.get("TUSALAMA_X_USER_IDS")
                    and os.environ.get("TUSALAMA_X_BEARER_TOKEN")
                ),
                "bluesky": bool(os.environ.get("TUSALAMA_BSKY_HANDLES")),
                "rss": bool(os.environ.get("TUSALAMA_RSS_FEEDS")),
            },
        )

    @app.get("/api/public-signals")
    def public_signals():
        now = clock.monotonic()
        if (
            public_signals_cache["payload"] is None
            or now - public_signals_cache["fetched_at"] >= 60
        ):
            public_signals_cache["payload"] = fetch_public_safety_signals()
            public_signals_cache["fetched_at"] = now
        return jsonify(public_signals_cache["payload"])

    @app.get("/reports/export")
    def export_report():
        incidents, _, _, _ = _filter_incidents(
            visible_incidents(), risk_assessment
        )
        export_format = request.args.get("format", "csv").lower()
        headers = [
            "Reference", "Date", "Time", "County", "Road", "Location",
            "Accident details", "Road user", "Fatalities", "Risk level",
            "Latitude", "Longitude",
        ]
        rows = [
            [
                item.get("cause_code", ""),
                item.get("date", ""),
                item.get("time", ""),
                item.get("county", ""),
                item.get("road", ""),
                item.get("place", ""),
                item.get("accident_info", ""),
                item.get("victim", ""),
                item.get("fatalities", ""),
                risk_assessment.get_risk_level(
                    risk_assessment.calculate_risk_score(item)
                ),
                item.get("latitude", ""),
                item.get("longitude", ""),
            ]
            for item in incidents
        ]
        safe_text_columns = {0, 3, 4, 5, 6, 7}

        def safe_spreadsheet_row(row):
            return [
                "'" + value if index in safe_text_columns
                and isinstance(value, str)
                and value.startswith(("=", "+", "-", "@"))
                else value
                for index, value in enumerate(row)
            ]

        if export_format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Incidents"
            sheet.append(headers)
            for row in rows:
                sheet.append(safe_spreadsheet_row(row))
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="166B4B")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                column_letter = column[0].column_letter
                sheet.column_dimensions[column_letter].width = min(
                    max(len(str(cell.value or "")) for cell in column) + 2, 38
                )
            output = io.BytesIO()
            workbook.save(output)
            response = make_response(output.getvalue())
            response.headers["Content-Type"] = (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
            response.headers["Content-Disposition"] = (
                "attachment; filename=tusalama-safety-report.xlsx"
            )
            return response
        if export_format == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.lib.units import mm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

            output = io.BytesIO()
            document = SimpleDocTemplate(
                output, pagesize=landscape(A4),
                leftMargin=12 * mm, rightMargin=12 * mm,
                topMargin=12 * mm, bottomMargin=12 * mm,
            )
            report_rows = [
                ["TUSALAMA ROAD SAFETY INCIDENT REPORT"],
                [f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} · {len(rows)} incidents"],
                headers,
            ] + [safe_spreadsheet_row(row) for row in rows]
            table = Table(
                report_rows,
                repeatRows=3,
                colWidths=[17, 20, 13, 21, 26, 25, 51, 22, 16, 19, 18, 18],
            )
            table.setStyle(TableStyle([
                ("SPAN", (0, 0), (-1, 0)),
                ("SPAN", (0, 1), (-1, 1)),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#166B4B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 1), 10),
                ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#EAF4EE")),
                ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
                ("FONTSIZE", (0, 2), (-1, -1), 6),
                ("GRID", (0, 2), (-1, -1), .25, colors.HexColor("#DDE5DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 3), (-1, -1), [colors.white, colors.HexColor("#F6F8F6")]),
            ]))
            document.build([table])
            response = make_response(output.getvalue())
            response.headers["Content-Type"] = "application/pdf"
            response.headers["Content-Disposition"] = (
                "attachment; filename=tusalama-safety-report.pdf"
            )
            return response
        if export_format != "csv":
            abort(400, description="Choose csv, xlsx, or pdf as the export format.")
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(["TUSALAMA ROAD SAFETY INCIDENT REPORT"])
        writer.writerow(["Generated", datetime.now().isoformat(timespec="minutes")])
        writer.writerow(["Total incidents", len(incidents)])
        writer.writerow([
            "Total fatalities",
            sum(int(item.get("fatalities") or 0) for item in incidents),
        ])
        writer.writerow([])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows(safe_spreadsheet_row(row) for row in rows)

        response = make_response(output.getvalue())
        response.headers["Content-Type"] = "text/csv; charset=utf-8"
        response.headers["Content-Disposition"] = (
            "attachment; filename=tusalama-safety-report.csv"
        )
        return response

    return app


app = create_app()


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1")
